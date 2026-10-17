import io
from xml.sax.saxutils import escape as xml_escape


def _section_break_indices(section_count: int, minimum_pages: int) -> set[int]:
    """Legacy compatibility helper retained for existing callers/tests."""
    return set()


def _page_groups(sections: list, target_pages: int | None = None) -> list[list]:
    """Keep sections in one natural flow; the layout engine owns pagination."""
    return [list(sections)] if sections else [[]]

from ...schemas import EditableDocument


def _pdf_text(value) -> str:
    """Escape report content before passing it to ReportLab Paragraph."""
    return xml_escape(str(value or ""))


def export_document(document: EditableDocument, output_format: str) -> tuple[bytes, str, str]:
    if output_format == "docx":
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        stream = io.BytesIO()
        doc = Document()
        columns = int(document.layout.get("columns", 1) or 1)
        if columns > 1:
            section = doc.sections[0]
            cols = section._sectPr.xpath("./w:cols")
            element = cols[0] if cols else OxmlElement("w:cols")
            if not cols:
                section._sectPr.append(element)
            element.set(qn("w:num"), str(columns))
        title = doc.add_heading(document.title, level=0)
        title.alignment = 1
        for index in range(0, len(document.authors), 3):
            row = document.authors[index:index + 3]
            paragraph = doc.add_paragraph()
            paragraph.alignment = 1
            paragraph.add_run("    ".join(f"{author.get('name', '[AUTHOR]')}\n{author.get('institution', '[INSTITUTION / AFFILIATION]')}" for author in row))
        date = doc.add_paragraph(document.date)
        date.alignment = 1
        for section in document.sections:
            doc.add_heading(section.heading, level=1)
            for paragraph in section.content.split("\n"):
                doc.add_paragraph(paragraph)
        doc.save(stream)
        return stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    if output_format == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        stream = io.BytesIO()
        width, height = letter
        columns = max(1, int(document.layout.get("columns", 1) or 1))
        gutter = 0.22 * inch
        frame_width = (width - 1.0 * inch - gutter * (columns - 1)) / columns
        frames = [Frame(0.5 * inch + i * (frame_width + gutter), 0.5 * inch,
                        frame_width, height - inch, id=f"col{i}", leftPadding=5, rightPadding=5,
                        topPadding=5, bottomPadding=5) for i in range(columns)]
        pdf = BaseDocTemplate(stream, pagesize=letter,
                              leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                              topMargin=0.5 * inch, bottomMargin=0.5 * inch)
        pdf.addPageTemplates([PageTemplate(id="report", frames=frames)])
        styles = getSampleStyleSheet()
        story = [Paragraph(_pdf_text(document.title), styles["Title"])]
        for index in range(0, len(document.authors), 3):
            row = document.authors[index:index + 3]
            author_text = " &nbsp;&nbsp;&nbsp; ".join(
                f"{_pdf_text(a.get('name', '[AUTHOR]'))}<br/>{_pdf_text(a.get('institution', '[INSTITUTION / AFFILIATION]'))}"
                for a in row
            )
            story.append(Paragraph(author_text, styles["BodyText"]))
        story.append(Paragraph(_pdf_text(document.date), styles["BodyText"]))
        for section in document.sections:
            story.append(Paragraph(_pdf_text(section.heading), styles["Heading2"]))
            story.extend(
                Paragraph(_pdf_text(paragraph), styles["BodyText"])
                for paragraph in section.content.split("\n") if paragraph.strip()
            )
        pdf.build(story)
        return stream.getvalue(), "application/pdf", "pdf"
    raise ValueError("Unsupported document format")
