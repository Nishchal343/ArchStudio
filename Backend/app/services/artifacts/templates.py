import io
import re
from pathlib import Path


# These are application-owned report formats. Uploaded files are not used for
# report generation; the source documents supplied for the product formats are
# represented here as section headings only.
MASTER_SECTIONS = [
    "Abstract", "Introduction", "Literature Review", "Research Gap & Motivation",
    "Proposed Methodology", "System Architecture", "Implementation Technology",
    "Comparison with Existing System", "Results and Discussion", "Future Scope",
    "Conclusion", "References",
]

PAPER_VARIANTS = {
    "report": {"columns": 1, "depth": "detailed", "page_target": "10 pages", "target_words": 4200, "output_tokens": 10500, "minimum_pages": 10, "page_count": 10},
    "ieee_report": {"columns": 2, "depth": "concise and technically dense", "page_target": "4 pages", "target_words": 1800, "output_tokens": 5000, "minimum_pages": 4, "page_count": 4},
}


def predefined_template(document_type: str) -> dict:
    template = PAPER_VARIANTS.get(document_type)
    if not template:
        raise ValueError("Unsupported predefined document type.")
    return {
        "layout": dict(template),
        "sections": [{"heading": heading, "level": 1} for heading in MASTER_SECTIONS],
    }


def analyze_template(filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        styles = [paragraph.style.name for paragraph in document.paragraphs if paragraph.text.strip()]
        columns = 1
        section = document.sections[0] if document.sections else None
        if section and section._sectPr.xpath("./w:cols/@w:num"):
            columns = int(section._sectPr.xpath("./w:cols/@w:num")[0])
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        lines = [line.strip() for page in reader.pages for line in (page.extract_text() or "").splitlines() if line.strip()]
        styles, columns = [], 1
        for page in reader.pages[:3]:
            x_positions = []
            def visitor(text, cm, tm, font, size):
                if text and text.strip():
                    x_positions.append(round(float(tm[4]), -1))
            try:
                page.extract_text(visitor_text=visitor)
            except Exception:
                x_positions = []
            if len(set(x_positions)) >= 2:
                ordered = sorted(set(x_positions))
                if ordered[-1] - ordered[0] > float(page.mediabox.width) * 0.35:
                    columns = 2
                    break
    else:
        raise ValueError("Template must be a PDF or DOCX file.")
    headings = []
    for index, line in enumerate(lines):
        numbered = re.match(r"^(\d+(?:\.\d+){0,3})[.)\s]+(.+)$", line)
        known = line.lower() in {"abstract", "keywords", "references", "conclusion", "appendix", "acknowledgment", "acknowledgements"}
        styled = index < len(styles) and styles[index].lower().startswith("heading")
        if numbered or known or styled:
            level = numbered.group(1).count(".") + 1 if numbered else (int(re.search(r"\d+", styles[index]).group()) if styled and re.search(r"\d+", styles[index]) else 1)
            headings.append({"heading": line, "level": level})
    if not headings:
        raise ValueError("No reliable section headings were detected in this template.")
    return {"filename": Path(filename).name, "layout": {"columns": columns}, "sections": headings}
